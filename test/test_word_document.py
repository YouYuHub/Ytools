"""Real DOCX fixtures: preserve table content/order and repair legacy text reads."""
import asyncio
import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from factory import file_factory
from factory import word_document as wd
from factory.agent_runtime import builtin_tools as bt
from factory.agent_runtime.chat_runtime import estimate_text_tokens
from memory import file_memory


def document_bytes(document):
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def sample_document():
    # Same failure as the model's report: table heading directly followed by the
    # closing paragraph when the old parser iterated document.paragraphs only.
    document = Document()
    document.add_heading("文档读取测试", 0)
    document.add_paragraph("第一段正文：中文、English、✅。")
    document.add_paragraph("列表第一项", style="List Bullet")
    document.add_heading("表格测试", 1)
    table = document.add_table(rows=3, cols=3)
    for row, values in zip(table.rows, [("姓名", "科目", "成绩"), ("张三", "语文", "88"), ("李四", "数学", "0")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("结尾段落：<尖括号> & “中文引号” —— 结束。")
    return document


class WordParserTests(unittest.TestCase):
    def test_report_reproduction_keeps_all_nine_cells_and_body_order(self):
        data = document_bytes(sample_document())
        text = file_factory.extract_text_from_bytes("测试.docx", data)
        self.assertIn("| 1 | 姓名 | 科目 | 成绩 |", text)
        self.assertIn("| 2 | 张三 | 语文 | 88 |", text)
        self.assertIn("| 3 | 李四 | 数学 | 0 |", text)
        self.assertLess(text.index("表格测试"), text.index("[表格 1]"))
        self.assertLess(text.index("| 3 |"), text.index("结尾段落"))
        self.assertIn("<尖括号> & “中文引号”", text)

    def test_multiple_tables_stay_between_their_surrounding_paragraphs(self):
        document = Document()
        for index in range(1, 3):
            document.add_paragraph(f"前文{index}")
            document.add_table(rows=1, cols=1).cell(0, 0).text = f"单元格{index}"
            document.add_paragraph(f"后文{index}")
        text = wd.extract_docx(document_bytes(document))
        locations = [text.index(value) for value in ["前文1", "单元格1", "后文1", "前文2", "单元格2", "后文2"]]
        self.assertEqual(locations, sorted(locations))

    def test_table_only_document_with_no_header_keeps_first_row(self):
        document = Document()
        document.add_table(rows=1, cols=2).cell(0, 0).text = "第一条数据"
        text = wd.extract_docx(document_bytes(document))
        self.assertIn("| 原始行号 | 列 1 | 列 2 |", text)
        self.assertIn("| 1 | 第一条数据 |  |", text)

    def test_cell_newlines_and_markdown_special_characters(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "<b>|`x`\n第二行"
        cell.add_paragraph("另一段")
        text = wd.extract_docx(document_bytes(document))
        self.assertIn("&lt;b&gt;\\|\\`x\\`<br>第二行<br>另一段", text)

    def test_rectangular_merge_preserves_anchor_once_and_marks_positions(self):
        document = Document()
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(1, 1)).text = "唯一合并内容"
        table.cell(2, 2).text = "末格"
        text = wd.extract_docx(document_bytes(document))
        self.assertEqual(text.count("唯一合并内容"), 1)
        self.assertEqual(text.count("[合并至 R1C1]"), 3)
        self.assertIn("| 3 |  |  | 末格 |", text)

    def test_separate_horizontal_and_vertical_merges(self):
        document = Document()
        table = document.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "横向"
        table.cell(1, 2).merge(table.cell(2, 2)).text = "纵向"
        text = wd.extract_docx(document_bytes(document))
        self.assertIn("[合并至 R1C1]", text)
        self.assertIn("[合并至 R2C3]", text)
        self.assertEqual(text.count("横向"), 1)
        self.assertEqual(text.count("纵向"), 1)

    def test_nested_table_has_reference_location_and_all_contents(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "嵌套之前"
        nested = cell.add_table(rows=2, cols=1)
        nested.cell(0, 0).text = "子表标题"
        nested.cell(1, 0).text = "子表数据"
        cell.add_paragraph("嵌套之后")
        text = wd.extract_docx(document_bytes(document))
        self.assertIn("嵌套之前<br>[嵌套表格 1.1]<br>嵌套之后", text)
        self.assertIn("[表格 1.1；位于表格 1 第 1 行第 1 列]", text)
        self.assertEqual(text.count("子表数据"), 1)

    def test_content_controls_wrapping_table_row_and_cell(self):
        document = Document()
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "内容控件中的表格数据"
        for element in [table.cell(0, 0)._tc, table.rows[0]._tr, table._tbl]:
            parent, index = element.getparent(), element.getparent().index(element)
            wrapper, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
            wrapper.append(content)
            content.append(element)
            parent.insert(index, wrapper)
        self.assertIn("| 1 | 内容控件中的表格数据 |", wd.extract_docx(document_bytes(document)))

    def test_omitted_leading_and_trailing_grid_cells_keep_alignment(self):
        document = Document()
        table = document.add_table(rows=1, cols=3)
        table.cell(0, 1).text = "中间"
        row = table.rows[0]._tr
        cells = list(row.findall(qn("w:tc")))
        row.remove(cells[0])
        row.remove(cells[2])
        for tag in ("w:gridBefore", "w:gridAfter"):
            prop = OxmlElement(tag)
            prop.set(qn("w:val"), "1")
            row.get_or_add_trPr().append(prop)
        self.assertIn("| 1 |  | 中间 |  |", wd.extract_docx(document_bytes(document)))

    def test_page_parser_matches_upload_parser_without_inventing_page_numbers(self):
        data = document_bytes(sample_document())
        text = file_factory.extract_text_from_bytes("x.docx", data)
        self.assertEqual(file_factory.extract_pages_from_bytes("x.docx", data), [{"content": text, "page_number": None}])

    def test_corrupt_docx_is_not_decoded_as_zip_text(self):
        for reader in (file_factory.extract_text_from_bytes, file_factory.extract_pages_from_bytes):
            with self.subTest(reader=reader.__name__), self.assertRaisesRegex(RuntimeError, "DOCX 解析失败"):
                reader("fake.docx", b"fake ordinary text")

    def test_cell_text_column_and_table_limits_are_explicit(self):
        data = document_bytes(sample_document())
        limits = [("MAX_WORD_TABLE_CELLS", 8), ("MAX_WORD_TEXT_CHARS", 5),
                  ("MAX_WORD_TABLE_COLUMNS", 2), ("MAX_WORD_TABLES", 0)]
        for name, limit in limits:
            with self.subTest(limit=name), patch.object(wd, name, limit), self.assertRaisesRegex(ValueError, "上限"):
                wd.extract_docx(data)

    def test_nested_table_depth_limit(self):
        document = Document()
        cell = document.add_table(rows=1, cols=1).cell(0, 0)
        cell.add_table(rows=1, cols=1).cell(0, 0).text = "子表"
        with patch.object(wd, "MAX_WORD_TABLE_DEPTH", 1), self.assertRaisesRegex(ValueError, "深度"):
            wd.extract_docx(document_bytes(document))

    def test_zip_expansion_limits_checked_before_document_is_loaded(self):
        data = document_bytes(sample_document())
        for name in ("MAX_WORD_EXPANDED_BYTES", "MAX_WORD_MEMBER_BYTES"):
            with self.subTest(limit=name), patch.object(wd, name, 100), patch.object(wd, "Document") as loader:
                with self.assertRaisesRegex(ValueError, "解压"):
                    wd.extract_docx(data)
                loader.assert_not_called()


class WordDocumentToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.folder = Path(self.tmp.name)
        self.root = patch.object(file_memory, "HISTORY_ROOT", self.folder)
        self.root.start()
        self.sid = "word_tool_test"
        self.data = document_bytes(sample_document())
        bt._document_text_cache.clear()

    def tearDown(self):
        asyncio.run(file_memory.cleanup_file_memory_manager(self.sid))
        self.root.stop()
        self.tmp.cleanup()
        bt._document_text_cache.clear()

    def local(self):
        path = self.folder / "测试.docx"
        path.write_bytes(self.data)
        return str(path)

    def legacy_record(self, content="只有旧正文", **fields):
        saved = file_memory.save_session_document(self.sid, "旧文档.docx", self.data)
        manager = asyncio.run(file_memory.get_file_memory_manager(self.sid))
        manager.add_file_memory({"filename": "旧文档.docx", "type": "docx", "content": content,
                                 "stored_name": saved["stored_name"], **fields})
        return file_memory._get_history_path(self.sid, "旧文档.docx"), file_memory.resolve_document_path(self.sid, saved["stored_name"])

    def test_auto_without_native_support_contains_word_table(self):
        result = bt.execute_read_document({"filename": self.local()}, self.sid, support_doc_types=[])
        self.assertEqual(result["mode"], "text")
        self.assertIn("| 2 | 张三 | 语文 | 88 |", result["content"])
        self.assertNotIn("native", result)
        self.assertFalse(result["source_truncated"])

    def test_explicit_text_with_native_support_still_parses_tables(self):
        result = bt.execute_read_document({"filename": self.local(), "mode": "text"}, self.sid, [".docx"])
        self.assertIn("| 1 | 姓名 | 科目 | 成绩 |", result["content"])
        self.assertNotIn("native", result)

    def test_legacy_untruncated_upload_recovers_tables_without_rewriting_record(self):
        record_path, _ = self.legacy_record()
        original_record = record_path.read_bytes()
        result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
        self.assertIn("| 3 | 李四 | 数学 | 0 |", result["content"])
        self.assertNotIn("parse_warning", result)
        self.assertEqual(record_path.read_bytes(), original_record)

    def test_legacy_empty_upload_can_recover_tables(self):
        self.legacy_record(content="", parse_failed=True)
        result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
        self.assertIn("张三", result["content"])
        self.assertNotIn("解析失败", result["message"])

    def test_word_table_mode_error_points_to_working_text_mode(self):
        result = bt.execute_read_document({"filename": self.local(), "mode": "table"}, self.sid)
        self.assertIn("Word 表格随正文返回为 Markdown", result["error"])
        self.assertIn("mode=text", result["error"])

    def test_truncated_record_restores_full_text_and_precise_continuation(self):
        self.legacy_record(content="文档", content_truncated=True, content_total_chars=9999)
        result = bt.execute_read_document({"filename": "旧文档.docx", "limit": 10}, self.sid)
        self.assertFalse(result["source_truncated"])
        self.assertEqual(result["next_offset"], 10)
        self.assertEqual(result["total_chars"], len(wd.extract_docx(self.data)))

    def test_missing_original_warns_that_legacy_text_may_lack_tables(self):
        _, raw_path = self.legacy_record()
        raw_path.unlink()
        result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
        self.assertEqual(result["content"], "只有旧正文")
        self.assertIn("旧记录可能缺失表格", result["parse_warning"])
        self.assertIn("重新上传", result["parse_warning"])

    def test_failed_original_reparse_exposes_warning_and_preserves_saved_text(self):
        _, raw_path = self.legacy_record()
        raw_path.write_bytes(b"not a zip")
        result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
        self.assertEqual(result["content"], "只有旧正文")
        self.assertIn("DOCX 解析失败", result["parse_warning"])

    def test_text_cache_reused_for_upload_pages_then_invalidated_on_change(self):
        _, raw_path = self.legacy_record()
        with patch.object(file_factory, "extract_text_from_bytes", wraps=file_factory.extract_text_from_bytes) as parser:
            bt.execute_read_document({"filename": "旧文档.docx", "limit": 5}, self.sid)
            bt.execute_read_document({"filename": "旧文档.docx", "offset": 5}, self.sid)
            self.assertEqual(parser.call_count, 1)
            document = sample_document()
            document.add_paragraph("新数据，文件已更新")
            raw_path.write_bytes(document_bytes(document))
            result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
            self.assertIn("新数据，文件已更新", result["content"])
            self.assertEqual(parser.call_count, 2)

    def test_small_model_budget_and_continuation_reassemble_the_whole_document(self):
        document = sample_document()
        table = document.add_table(rows=70, cols=2)
        for index, row in enumerate(table.rows):
            row.cells[0].text = str(index)
            row.cells[1].text = "预算测试数据" * 3
        data = document_bytes(document)
        path = self.folder / "long.docx"
        path.write_bytes(data)
        offset, pieces = 0, []
        while True:
            result = bt.execute_read_document({"filename": str(path), "offset": offset, "limit": 20000,
                                              "_max_tokens": 400}, self.sid)
            self.assertNotIn("error", result)
            self.assertLessEqual(estimate_text_tokens(result), 400)
            self.assertEqual(result["start_char"], offset)
            pieces.append(result["content"])
            if not result["has_more"]:
                break
            self.assertGreater(result["next_offset"], offset)
            offset = result["next_offset"]
        self.assertEqual("".join(pieces), wd.extract_docx(data))

    def test_legacy_raw_text_parse_is_bounded_separately_from_native_upload(self):
        self.legacy_record()
        with patch.object(bt, "_READ_DOCUMENT_LOCAL_MAX_BYTES", 100):
            result = bt.execute_read_document({"filename": "旧文档.docx"}, self.sid)
        self.assertIn("parse_warning", result)
        self.assertIn("上限", result["parse_warning"])
        self.assertEqual(result["content"], "只有旧正文")


if __name__ == "__main__":
    unittest.main()
