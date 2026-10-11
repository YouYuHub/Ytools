"""Bounded spreadsheet parsing and Markdown rendering shared by document tools."""
import csv
import io
import itertools
import math
import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from threading import Lock

TABLE_EXTENSIONS = frozenset({".csv", ".tsv", ".xls", ".xlsx"})
MAX_TABLE_ROWS = 50000
MAX_TABLE_COLUMNS = 1024
MAX_TABLE_CELLS = 200000
MAX_TABLE_CHARS = 4 * 1024 * 1024
MAX_XLSX_EXPANDED_BYTES = 128 * 1024 * 1024
_csv_lock = Lock()


@dataclass
class TableSheet:
    name: str
    rows: list[list[str]]


@dataclass
class TableDocument:
    sheets: list[TableSheet]
    chars: int = 0
    cells: int = 0
    row_count: int = 0
    formula_fallback: bool = False

    def append_sheet(self, name, rows):
        collected = []
        for row in rows:
            if len(collected) >= MAX_TABLE_ROWS or len(row) > MAX_TABLE_COLUMNS:
                raise ValueError("表格超过解析行列上限（每表 50000 行、1024 列）；请拆分文件或用终端筛选")
            values = [cell_text(value) for value in row]
            self.row_count += 1
            self.cells += len(values)
            self.chars += sum(map(len, values))
            if self.cells > MAX_TABLE_CELLS or self.row_count > MAX_TABLE_CELLS or self.chars > MAX_TABLE_CHARS:
                raise ValueError("表格超过解析预算（20 万行/单元格或 4 Mi 字符）；请拆分文件或用终端筛选")
            collected.append(values)
        self.sheets.append(TableSheet(str(name), collected))


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isfinite(value) and value.is_integer():
            return str(int(value))
    return str(value)


def column_name(index):
    name = ""
    while index:
        index, digit = divmod(index - 1, 26)
        name = chr(65 + digit) + name
    return name


def markdown_cell(value):
    # Escape separators/HTML and keep a multiline cell inside a single table row.
    return (str(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            .replace("\\", "\\\\").replace("|", "\\|").replace("`", "\\`")
            .replace("\r\n", "\n").replace("\r", "\n").replace("\n", "<br>").replace("\t", " "))


def _xlsx_sheets(data, document):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("XLSX 解析需要 openpyxl；请安装 requirements.txt 的文档解析依赖") from exc
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sizes = [item.file_size for item in archive.infolist()]
        if sum(sizes) > MAX_XLSX_EXPANDED_BYTES or any(size > 64 * 1024 * 1024 for size in sizes):
            raise ValueError("XLSX 解压后过大，拒绝全量解析；请拆分文件")
    # Stream both forms so uncached formula cells remain visible instead of becoming empty.
    formulas = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            if len(formulas.sheetnames) > 100:
                raise ValueError("工作簿超过 100 个工作表，请拆分文件")
            for sheet in formulas:
                values = cached[sheet.title]
                # Dimensions supplied by exporters can be incorrect. Discover the actual rows.
                sheet.reset_dimensions()
                values.reset_dimensions()
                def rows():
                    for actual, calculated in itertools.zip_longest(
                            sheet.iter_rows(values_only=False), values.iter_rows(values_only=True), fillvalue=()):
                        row = []
                        for index, cell in enumerate(actual):
                            value = cell.value
                            if cell.data_type == "f":
                                cached_value = calculated[index] if index < len(calculated) else None
                                if cached_value is not None:
                                    value = cached_value
                                else:
                                    document.formula_fallback = True
                            row.append(value)
                        yield row
                document.append_sheet(sheet.title, rows())
        finally:
            cached.close()
    finally:
        formulas.close()


def _xls_sheets(data, document):
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("XLS 解析需要 xlrd；请安装 xlrd，或将文件另存为 XLSX/CSV") from exc
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        if workbook.nsheets > 100:
            raise ValueError("工作簿超过 100 个工作表，请拆分文件")
        for number in range(workbook.nsheets):
            sheet = workbook.sheet_by_index(number)
            def rows():
                for row_index in range(sheet.nrows):
                    row = []
                    for cell in sheet.row(row_index):
                        value = cell.value
                        if cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate_as_datetime(value, workbook.datemode)
                        elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                            value = bool(value)
                        elif cell.ctype == xlrd.XL_CELL_ERROR:
                            value = xlrd.error_text_from_code.get(value, f"#ERROR({value})")
                        row.append(value)
                    yield row
            document.append_sheet(sheet.name, rows())
            workbook.unload_sheet(number)
    finally:
        workbook.release_resources()


def parse_table(filename, data):
    extension = Path(filename).suffix.lower()
    document = TableDocument([])
    if extension in {".csv", ".tsv"}:
        from util.text_encoding import decode_text
        text, _ = decode_text(data, candidates=("utf-8-sig", "gb18030", "big5"))
        delimiter = "\t" if extension == ".tsv" else ","
        if extension == ".csv":
            try:
                delimiter = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|").delimiter
            except csv.Error:
                pass
        with _csv_lock:
            old_limit = csv.field_size_limit()
            try:
                csv.field_size_limit(MAX_TABLE_CHARS)
                document.append_sheet("CSV" if extension == ".csv" else "TSV",
                                      csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
            finally:
                csv.field_size_limit(old_limit)
    elif extension == ".xlsx":
        _xlsx_sheets(data, document)
    elif extension == ".xls":
        _xls_sheets(data, document)
    else:
        raise ValueError("表格模式仅支持 CSV/TSV/XLS/XLSX 文件")
    return document


def sheet_headers(sheet, header_row):
    if header_row < 0 or header_row > len(sheet.rows):
        raise ValueError(f"header_row 超出工作表范围（共 {len(sheet.rows)} 行）")
    width = max(map(len, sheet.rows), default=0)
    source = sheet.rows[header_row - 1] if header_row else []
    return [source[index] if index < len(source) and source[index] else f"列 {column_name(index + 1)}"
            for index in range(width)]


def table_page(document, sheet_name=None, *, header_row=1, start_row=None, row_limit=100, max_chars=10000):
    names = [sheet.name for sheet in document.sheets]
    if not names:
        raise ValueError("工作簿没有可读取的工作表")
    if sheet_name is not None and sheet_name not in names:
        raise ValueError(f"工作表 {sheet_name!r} 不存在；可用工作表：{', '.join(names)}")
    number = names.index(sheet_name) if sheet_name is not None else 0
    sheet = document.sheets[number]
    if not sheet.rows:
        headers, header_row = [], 0
    else:
        headers = sheet_headers(sheet, header_row)
    first = header_row + 1 if header_row else 1
    start_row = first if start_row is None else start_row
    if start_row < first:
        raise ValueError(f"start_row 必须 >= {first}（数据从表头之后开始）；无表头请设 header_row=0")
    if start_row > len(sheet.rows) + 1:
        raise ValueError(f"start_row 超出工作表范围（共 {len(sheet.rows)} 行）")
    prefix = "### 工作表：" + markdown_cell(sheet.name) + "\n\n"
    if headers:
        prefix += "| 原始行号 | " + " | ".join(map(markdown_cell, headers)) + " |\n"
        prefix += "| --- | " + " | ".join("---" for _ in headers) + " |\n"
    if len(prefix) > max_chars:
        raise ValueError("表头超过本次字符预算；请增大 limit，或使用 mode=text 按字符读取")
    lines = [prefix]
    length = len(prefix)
    count = 0
    for index in range(start_row - 1, min(len(sheet.rows), start_row - 1 + row_limit)):
        cells = sheet.rows[index]
        line = f"| {index + 1} | " + " | ".join(markdown_cell(cells[col] if col < len(cells) else "")
                                                for col in range(len(headers))) + " |\n"
        if length + len(line) > max_chars:
            if count == 0:
                raise ValueError("单行超过本次字符预算；请增大 limit，或使用 mode=text 按字符读取完整单元格")
            break
        lines.append(line)
        length += len(line)
        count += 1
    end = start_row + count - 1
    next_row = end + 1 if end < len(sheet.rows) else None
    next_sheet = names[number + 1] if number + 1 < len(names) else None
    if not sheet.rows or start_row > len(sheet.rows):
        next_row = None
    result = {"format": "markdown_table", "sheet": sheet.name, "sheets": names,
              "header_row": header_row, "headers": headers, "total_rows": len(sheet.rows),
              "start_row": start_row, "end_row": end, "rows_returned": count,
              "content": "".join(lines), "has_more": bool(next_row or next_sheet)}
    if next_row:
        result["next_row"] = next_row
    elif next_sheet:
        result["next_sheet"] = next_sheet
    if document.formula_fallback:
        result["formula_note"] = "公式优先读取文件保存的缓存值；无缓存时保留公式表达式，未执行计算"
    return result


def markdown_sheets(document):
    """Full parsed text used by upload manifests and legacy character pagination."""
    for sheet in document.sheets:
        headers = sheet_headers(sheet, 1) if sheet.rows else []
        lines = ["### 工作表：" + markdown_cell(sheet.name), ""]
        if headers:
            lines += ["| 原始行号 | " + " | ".join(map(markdown_cell, headers)) + " |",
                      "| --- | " + " | ".join("---" for _ in headers) + " |"]
            for number, row in enumerate(sheet.rows[1:], 2):
                lines.append(f"| {number} | " + " | ".join(markdown_cell(row[col] if col < len(row) else "")
                                                           for col in range(len(headers))) + " |")
        yield "\n".join(lines)
