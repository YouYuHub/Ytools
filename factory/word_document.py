"""DOCX text fallback: ordered paragraphs and readable tables, without layout rendering."""
import io
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from factory.table_document import markdown_cell

MAX_WORD_TEXT_CHARS = 4 * 1024 * 1024
MAX_WORD_TABLE_CELLS = 200000
MAX_WORD_TABLES = 1000
MAX_WORD_TABLE_COLUMNS = 1024
MAX_WORD_TABLE_DEPTH = 16
MAX_WORD_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_WORD_MEMBER_BYTES = 64 * 1024 * 1024


def _elements(container, tags, depth=0):
    """Walk direct content in order, unwrapping Word content controls/revisions."""
    if depth > 32:
        raise ValueError("DOCX 内容控件嵌套过深；请简化文档")
    for child in container:
        if child.tag in tags:
            yield child
        elif child.tag == qn("w:sdt"):
            content = child.find(qn("w:sdtContent"))
            if content is not None:
                yield from _elements(content, tags, depth + 1)
        elif child.tag in (qn("w:customXml"), qn("w:ins"), qn("w:moveTo")):
            yield from _elements(child, tags, depth + 1)


def _blocks(container):
    return _elements(container, (qn("w:p"), qn("w:tbl")))


def _number(element, name, default=0):
    found = element.find(qn(name)) if element is not None else None
    return max(0, int(found.get(qn("w:val"), default))) if found is not None else default


class _Renderer:
    def __init__(self):
        self.text_chars = 0
        self.cells = 0
        self.tables = 0

    def paragraph(self, element, parent):
        text = Paragraph(element, parent).text
        text = "".join(char for char in text if ord(char) >= 32 or char in "\n\r\t").strip()
        self.text_chars += len(text)
        if self.text_chars > MAX_WORD_TEXT_CHARS:
            raise ValueError("DOCX 文本超过解析上限；请拆分文档，不会静默丢弃表格或段落")
        return text

    def table(self, element, parent, label, location="", depth=0):
        self.tables += 1
        if self.tables > MAX_WORD_TABLES or depth >= MAX_WORD_TABLE_DEPTH:
            raise ValueError("DOCX 表格数量或嵌套深度超过解析上限；请拆分文档")
        table = Table(element, parent)
        rows = list(_elements(element, (qn("w:tr"),)))
        widths = []
        for row in rows:
            properties = row.find(qn("w:trPr"))
            widths.append(_number(properties, "w:gridBefore") + _number(properties, "w:gridAfter") +
                          sum(max(1, _number(cell.find(qn("w:tcPr")), "w:gridSpan", 1))
                              for cell in _elements(row, (qn("w:tc"),))))
        grid = element.find(qn("w:tblGrid"))
        columns = max([len(grid) if grid is not None else 0, *widths])
        self.cells += len(rows) * columns
        if columns > MAX_WORD_TABLE_COLUMNS or self.cells > MAX_WORD_TABLE_CELLS:
            raise ValueError("DOCX 表格列数或单元格数量超过解析上限；请拆分文档")
        caption = f"[表格 {label}" + (f"；{location}" if location else "") + "]"
        if not rows or not columns:
            return caption + "\n（空表格）"

        rendered_rows, nested_tables, previous_merges = [], [], {}
        nested_number = 0
        for row_number, row in enumerate(rows, 1):
            properties = row.find(qn("w:trPr"))
            column = _number(properties, "w:gridBefore")
            values, current_merges = [""] * columns, {}
            for cell_element in _elements(row, (qn("w:tc"),)):
                cell = _Cell(cell_element, table)
                parts = []
                for block in _blocks(cell_element):
                    if block.tag == qn("w:p"):
                        text = self.paragraph(block, cell)
                        if text:
                            parts.append(text)
                    else:
                        nested_number += 1
                        nested_label = f"{label}.{nested_number}"
                        parts.append(f"[嵌套表格 {nested_label}]")
                        nested_tables.append(self.table(
                            block, cell, nested_label,
                            f"位于表格 {label} 第 {row_number} 行第 {column + 1} 列", depth + 1))
                cell_properties = cell_element.find(qn("w:tcPr"))
                span = max(1, _number(cell_properties, "w:gridSpan", 1))
                merge = cell_properties.find(qn("w:vMerge")) if cell_properties is not None else None
                continuation = merge is not None and merge.get(qn("w:val"), "continue") == "continue"
                origin = previous_merges.get(column) if continuation else (row_number, column + 1)
                text = "\n".join(parts)
                for position in range(column, column + span):
                    if continuation or position > column:
                        marker = (f"[合并至 R{origin[0]}C{origin[1]}]" if origin else "[纵向合并，起始单元格缺失]")
                        values[position] = (text + "\n" + marker) if position == column and text else marker
                    else:
                        values[position] = text
                    if merge is not None and origin is not None:
                        current_merges[position] = origin
                column += span
            previous_merges = current_merges
            rendered_rows.append([str(row_number), *values])

        # Word tables can start with data, titles or merged cells. Neutral column labels
        # preserve every source row instead of silently treating the first row as a header.
        headers = ["原始行号", *(f"列 {index}" for index in range(1, columns + 1))]
        lines = [caption, "| " + " | ".join(headers) + " |",
                 "| " + " | ".join("---" for _ in headers) + " |"]
        lines.extend("| " + " | ".join(markdown_cell(value) for value in row) + " |" for row in rendered_rows)
        return "\n\n".join(["\n".join(lines), *nested_tables])

    def body(self, document):
        blocks, table_number = [], 0
        for block in _blocks(document.element.body):
            if block.tag == qn("w:p"):
                text = self.paragraph(block, document)
            else:
                table_number += 1
                text = self.table(block, document, str(table_number))
            if text:
                blocks.append(text)
        text = "\n\n".join(blocks)
        if len(text) > MAX_WORD_TEXT_CHARS:
            raise ValueError("DOCX 格式化文本超过解析上限；请拆分文档")
        return text


def extract_docx(data: bytes) -> str:
    """Validate archive size before python-docx expands it, then preserve body order."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        sizes = [item.file_size for item in archive.infolist()]
        if sum(sizes) > MAX_WORD_EXPANDED_BYTES or any(size > MAX_WORD_MEMBER_BYTES for size in sizes):
            raise ValueError("DOCX 解压后过大，拒绝全量解析；请拆分文件")
    return _Renderer().body(Document(io.BytesIO(data)))
